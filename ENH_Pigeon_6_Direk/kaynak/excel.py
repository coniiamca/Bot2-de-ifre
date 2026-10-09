# -*- coding: utf-8 -*-
"""Formullu Excel: kesif (birim fiyat girisli), tepe kuvveti, iletken mekanigi, temel/beton, proje esaslari."""
import json
import math
import os
import sys

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.page import PageMargins

HERE = os.path.dirname(os.path.abspath(__file__))
sys.path.insert(0, HERE)
import kesif as KS  # noqa: E402

S = json.load(open(os.path.join(HERE, "hesap_sonuclari.json"), encoding="utf-8"))
OUT = os.path.join(os.path.dirname(HERE), "ENH_Pigeon_AER_6_Direk_Kesif_ve_Hesaplar.xlsx")

ince = Side(style="thin", color="808080")
kalin = Side(style="medium", color="000000")
CER = Border(left=ince, right=ince, top=ince, bottom=ince)
BAS_FONT = Font(name="Arial", bold=True, size=10, color="FFFFFF")
BAS_FILL = PatternFill("solid", fgColor="1F4E78")
BOL_FILL = PatternFill("solid", fgColor="D9E1F2")
GIR_FILL = PatternFill("solid", fgColor="FFF2CC")       # kullanici girisi
TOP_FILL = PatternFill("solid", fgColor="E2EFDA")
F = Font(name="Arial", size=9)
FB = Font(name="Arial", size=9, bold=True)
ORTA = Alignment(horizontal="center", vertical="center", wrap_text=True)
SOL = Alignment(horizontal="left", vertical="center", wrap_text=True)
SAG = Alignment(horizontal="right", vertical="center")

DIR = ["D1", "D2", "D3", "D4", "D5", "D6"]
D = S["direkler"]
T = S["temel"]
IL = S["iletkenler"]


def baslik_satiri(ws, r, basliklar, genislik=None):
    for c, h in enumerate(basliklar, 1):
        cell = ws.cell(row=r, column=c, value=h)
        cell.font, cell.fill, cell.alignment, cell.border = BAS_FONT, BAS_FILL, ORTA, CER
    ws.row_dimensions[r].height = 32
    if genislik:
        for c, w in enumerate(genislik, 1):
            ws.column_dimensions[get_column_letter(c)].width = w


def yaz(ws, r, c, v, fmt=None, font=F, al=None, fill=None, border=CER):
    cell = ws.cell(row=r, column=c, value=v)
    cell.font = font
    cell.alignment = al or (SOL if isinstance(v, str) and len(v) > 12 else ORTA)
    if fmt:
        cell.number_format = fmt
    if fill:
        cell.fill = fill
    if border:
        cell.border = border
    return cell


def ust_bilgi(ws, baslik, alt, ncol):
    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=ncol)
    ws.merge_cells(start_row=2, start_column=1, end_row=2, end_column=ncol)
    c = ws.cell(row=1, column=1, value=baslik)
    c.font = Font(name="Arial", bold=True, size=13)
    c = ws.cell(row=2, column=1, value=alt)
    c.font = Font(name="Arial", italic=True, size=9)
    ws.row_dimensions[1].height = 22


def sayfa_ayari(ws, yatay=True):
    ws.page_setup.orientation = "landscape" if yatay else "portrait"
    ws.page_setup.paperSize = ws.PAPERSIZE_A3
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.page_margins = PageMargins(left=0.4, right=0.4, top=0.5, bottom=0.5)


wb = Workbook()

# ================================================================== KESIF
ws = wb.active
ws.title = "KEŞİF"
KOL = KS.KOLON
bas = ["S.No", "TEDAŞ Poz No\n(siz girin)", "Malzeme / İş Tanımı", "Birim"] + KOL + \
      ["TOPLAM\nMİKTAR", "BİRİM FİYAT\n(TL, siz girin)", "TUTAR\n(TL)", "Açıklama"]
gen = [6, 14, 78, 8] + [8] * len(KOL) + [11, 15, 16, 30]
ust_bilgi(ws, "34,5 kV 3x3/0 AWG (Pigeon) OG + 0,4 kV AER 3x95+70 mm² AG MÜŞTEREK DEMİR DİREKLİ HAVA HATTI - "
              "MALZEME VE İŞ KEŞFİ",
          "6 direk, III. buz yükü bölgesi. Sarı hücreler (TEDAŞ poz no, birim fiyat) kullanıcı girişidir; "
          "tutar, ara toplam ve genel toplam formüllüdür. M: mevcut direkte bağlantı, Hat: güzergâh boyunca.",
          len(bas))
R0 = 4
baslik_satiri(ws, R0, bas, gen)
ws.freeze_panes = ws.cell(row=R0 + 1, column=4)
bolumler, ozet = KS.kesif()
r = R0 + 1
n = 0
c_top = 4 + len(KOL) + 1
c_bf, c_tut, c_ac = c_top + 1, c_top + 2, c_top + 3
L_top, L_bf, L_tut = get_column_letter(c_top), get_column_letter(c_bf), get_column_letter(c_tut)
ara_toplamlar = []
for bl in bolumler:
    ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=c_ac)
    yaz(ws, r, 1, bl["baslik"], font=FB, fill=BOL_FILL, al=SOL)
    r += 1
    r_bas = r
    for it in bl["kalemler"]:
        n += 1
        yaz(ws, r, 1, n)
        yaz(ws, r, 2, None, fill=GIR_FILL)
        yaz(ws, r, 3, it["tanim"], al=SOL)
        yaz(ws, r, 4, it["birim"])
        for j, k in enumerate(KOL):
            v = it["miktar"][k]
            yaz(ws, r, 5 + j, v if v else None, fmt="#,##0.###")
        yaz(ws, r, c_top, f"=SUM({get_column_letter(5)}{r}:{get_column_letter(4 + len(KOL))}{r})", fmt="#,##0.###",
            font=FB)
        yaz(ws, r, c_bf, None, fmt='#,##0.00', fill=GIR_FILL)
        yaz(ws, r, c_tut, f'=IF({L_bf}{r}="","",{L_top}{r}*{L_bf}{r})', fmt='#,##0.00')
        yaz(ws, r, c_ac, it.get("not") or None, al=SOL)
        ws.row_dimensions[r].height = 26 if len(it["tanim"]) > 85 else 15
        r += 1
    ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=c_bf)
    yaz(ws, r, 1, f"{bl['baslik']} - ARA TOPLAM", font=FB, al=Alignment(horizontal="right"), fill=TOP_FILL)
    yaz(ws, r, c_tut, f"=SUM({L_tut}{r_bas}:{L_tut}{r - 1})", fmt='#,##0.00', font=FB, fill=TOP_FILL)
    yaz(ws, r, c_ac, None, fill=TOP_FILL)
    ara_toplamlar.append(f"{L_tut}{r}")
    r += 2
for lab, f_ in (("GENEL TOPLAM (KDV HARİÇ)", "=" + "+".join(ara_toplamlar)),
                ("KDV (%20)", f"={L_tut}{r}*0.20"),
                ("GENEL TOPLAM (KDV DAHİL)", f"={L_tut}{r}+{L_tut}{r + 1}")):
    ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=c_bf)
    yaz(ws, r, 1, lab, font=Font(name="Arial", bold=True, size=11), al=Alignment(horizontal="right"), fill=TOP_FILL)
    yaz(ws, r, c_tut, f_, fmt='#,##0.00', font=Font(name="Arial", bold=True, size=11), fill=TOP_FILL)
    r += 1
r += 1
notlar = [
    "NOTLAR:",
    "1. Miktarlar hesap_sonuclari.json'daki proje hesaplarından otomatik üretilmiştir (kaynak/kesif.py).",
    f"2. Pigeon iletken: 3 faz x {ozet['hat_og']:.0f} m güzergâh x 1,02 (sehim/kesim payı) + atlamalar = {ozet['og_iletken']:.1f} m.",
    f"3. AER 3x95+70: {ozet['hat_ag']:.0f} m x 1,03 + köşe by-pass ve uç payları = {ozet['ag_kablo']:.1f} m.",
    f"4. Beton (C20/25) toplamı {ozet['beton']:.3f} m³, kazı {ozet['kazi']:.3f} m³, kalıp {ozet['kalip']:.2f} m² (net; "
    "zayiat payı birim fiyat analizinde değerlendirilir).",
    f"5. Demir direk ağırlıkları TEDAŞ müşterek direk tablosundan (III. bölge, '' tipler): toplam {ozet['demir_kg']:.2f} kg; "
    "potans/konsol ağırlıkları hariçtir.",
    "6. Birim fiyat ve poz numaraları güncel TEDAŞ birim fiyat kitabından girilecektir; KDV oranı gerekirse değiştirilebilir.",
    "7. M kolonu: mevcut 34,5 kV hattın M direğinde yeni branşman bağlantısı için gerekli malzemelerdir.",
    "8. AER 3x95+70 taşıyıcı nötr kesiti (70/95 mm²) dağıtım şirketinin güncel malzeme listesine göre teyit edilmelidir.",
]
for t_ in notlar:
    ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=c_ac)
    ws.cell(row=r, column=1, value=t_).font = FB if t_ == "NOTLAR:" else F
    r += 1
sayfa_ayari(ws)
ws.print_title_rows = f"{R0}:{R0}"

# ============================================================ TEPE KUVVETI
ws = wb.create_sheet("TEPE KUVVETİ")
ust_bilgi(ws, "DİREK TEPE KUVVETİ HESAP CETVELİ (EKATY Md.48-49)",
          "Kuvvetler OG bağlantı seviyesine (H = 9,55 m) moment eşitliği ile indirgenmiştir: F = Σ Fi·hi/H. "
          "Köşe/nihayet bileşkeleri vektöreldir. Birim: kg (kgf).", 18)
bas = ["Direk", "Görevi", "Seçilen tip", "İç açı β (°)", "n = 2cos(β/2)", "a1 (m)", "a2 (m)", "aw (m)",
       "T_OG (kg/iletken)", "T_AG (kg)", "h_AG (m)", "h_AG/H", "Hesap tepe kuvveti (kg)", "Belirleyici varsayım",
       "Katalog tepe kuvveti (kg)", "En büyük kullanım", "Ağırlık (kg)", "Sonuç"]
baslik_satiri(ws, 4, bas, [7, 30, 10, 9, 10, 8, 8, 8, 11, 10, 9, 9, 14, 34, 13, 11, 11, 10])
kom = {"D1": (35, 45), "D2": (45, 50), "D3": (50, 50), "D4": (50, 50), "D5": (50, 45), "D6": (45, None)}
KISA2 = {"D1": "Seksiyoner + köşe durdurucu", "D2": "Köşe durdurucu", "D3": "Taşıyıcı", "D4": "Taşıyıcı",
         "D5": "Köşe durdurucu", "D6": "Nihayet"}
r = 5
for d in DIR:
    dd = D[d]
    beta = dd["ic_aci"]
    yaz(ws, r, 1, d, font=FB)
    yaz(ws, r, 2, KISA2[d], al=SOL)
    yaz(ws, r, 3, dd["tip"], font=FB)
    yaz(ws, r, 4, beta if beta else "-")
    if beta and beta < 179:
        yaz(ws, r, 5, f"=2*COS(RADIANS(D{r}/2))", fmt="0.000")
    elif beta:
        yaz(ws, r, 5, 0, fmt="0.000")
    else:
        yaz(ws, r, 5, 1, fmt="0.000")
    yaz(ws, r, 6, kom[d][0])
    yaz(ws, r, 7, kom[d][1] if kom[d][1] else "-")
    yaz(ws, r, 8, dd["aw_og"], fmt="0.0")
    yaz(ws, r, 9, IL["OG"]["Tmax"], fmt="0.0")
    yaz(ws, r, 10, IL["AG"]["Tmax"], fmt="0.0")
    yaz(ws, r, 11, dd["h_ag"], fmt="0.00")
    yaz(ws, r, 12, f"=K{r}/9.55", fmt="0.000")
    if dd["gorev"].startswith("Taşıyıcı"):
        yaz(ws, r, 13, f"{dd['varsayim']['V1 (hatta dik rüzgâr)']['F_hesap']:.1f} / "
                       f"{dd['varsayim']['V3 (iletken kopması)']['F_hesap']:.1f}")
        kat = S["katalog"][dd["tip"]]["tepe"]
        yaz(ws, r, 15, f"{kat[2]} / {kat[1]}")
    else:
        yaz(ws, r, 13, round(dd["F_tepe"], 1), fmt="#,##0.0", font=FB)
        yaz(ws, r, 15, dd["kapasite"], fmt="#,##0")
    yaz(ws, r, 14, dd["belirleyici"], al=SOL)
    yaz(ws, r, 16, dd["max_kullanim"], fmt="0%")
    yaz(ws, r, 17, dd["agirlik"], fmt="#,##0.00")
    yaz(ws, r, 18, "UYGUN" if dd["max_kullanim"] <= 1 else "UYGUN DEĞİL", font=FB)
    r += 1
yaz(ws, r, 16, "Toplam", font=FB)
yaz(ws, r, 17, f"=SUM(Q5:Q{r - 1})", fmt="#,##0.00", font=FB)
r += 2
ws.cell(row=r, column=1, value="Taşıyıcılarda: hatta dik rüzgâr (güçlü eksen, rüzgârlı katalog değeri) / iletken kopması "
        "(zayıf eksen; OG mesnet izolatörlü T/5, AG demet T/3).").font = F
r += 2
ws.cell(row=r, column=1, value="YÜKLENME VARSAYIMLARI AYRINTISI").font = Font(name="Arial", bold=True, size=11)
r += 1
bas2 = ["Direk", "Tip", "Varsayım", "Açıklama / bileşenler", "Eksen", "Kapasite türü", "F hesap (kg)",
        "Direk rüzgârı (tepe, kg)", "Travers/izolatör rüzgârı (kg)", "Kapasite (kg)", "Kullanım", "Sonuç"]
for c, h in enumerate(bas2, 1):
    cell = ws.cell(row=r, column=c, value=h)
    cell.font, cell.fill, cell.alignment, cell.border = BAS_FONT, BAS_FILL, ORTA, CER
r += 1
for d in DIR:
    dd = D[d]
    for vn, v in dd["varsayim"].items():
        yaz(ws, r, 1, d)
        yaz(ws, r, 2, dd["tip"])
        yaz(ws, r, 3, vn, al=SOL)
        yaz(ws, r, 4, v["aciklama"], al=SOL)
        yaz(ws, r, 5, {"guclu": "güçlü", "zayif": "zayıf", "her": "her yön"}[v["eksen"]])
        yaz(ws, r, 6, "rüzgârsız" if S["katalog"][dd["tip"]]["cins"] == "K" else
            v["kap"].replace("ruzgarli", "rüzgârlı").replace("ruzgarsiz", "rüzgârsız"))
        yaz(ws, r, 7, round(v["F_hesap"], 1), fmt="#,##0.0")
        yaz(ws, r, 8, round(v.get("W_direk_tepe", 0), 1) or None, fmt="0.0")
        yaz(ws, r, 9, round(v.get("W_travers", 0), 1) or None, fmt="0.0")
        yaz(ws, r, 10, v["kapasite"], fmt="#,##0")
        yaz(ws, r, 11, f"=G{r}/J{r}", fmt="0%")
        yaz(ws, r, 12, f'=IF(K{r}<=1,"UYGUN","UYGUN DEĞİL")', font=FB)
        r += 1
sayfa_ayari(ws)

# ======================================================= ILETKEN MEKANIGI
ws = wb.create_sheet("İLETKEN MEKANİĞİ")
ust_bilgi(ws, "İLETKEN GERİLME - SEHİM CETVELİ (durum değişim denklemi)",
          "σ2²·(σ2 - [σ1 - E·a²·γ1²/(24·σ1²) - α·E·(t2 - t1)]) = E·a²·γ2²/24 ; γ = w/S. Germe bölümlerinde eşdeğer açıklık "
          "a_e = √(Σa³/Σa). En büyük çekme belirleyici durumda Tmax'a eşitlenmiştir.", 16)
bas = ["Devre", "Germe bölümü", "Eşdeğer açıklık (m)", "Belirleyici durum", "T -25°C (kg)", "T -5°C+buz (kg)",
       "T +5°C+rüzgâr (kg)", "T +15°C EDS (kg)", "EDS / kopma", "T +40°C (kg)", "Açıklık", "a (m)",
       "Sehim -5°C+buz (m)", "Sehim +40°C (m)", "Sehim -25°C (m)", "Sehim +15°C (m)"]
baslik_satiri(ws, 4, bas, [7, 12, 11, 12, 11, 11, 12, 12, 10, 11, 10, 8, 12, 12, 12, 12])
r = 5
for dev in ("OG", "AG"):
    for bk, m in S["mekanik"][dev].items():
        for j, (ac, sh) in enumerate(m["sehim"].items()):
            if j == 0:
                yaz(ws, r, 1, dev, font=FB)
                yaz(ws, r, 2, bk)
                yaz(ws, r, 3, round(m["esdeger_aciklik"], 2), fmt="0.0")
                yaz(ws, r, 4, {"buz": "-5°C + buz", "min": "-25°C", "ruz": "+5°C + rüzgâr"}[m["belirleyici"]])
                for c, k in ((5, "min"), (6, "buz"), (7, "ruz"), (8, "eds"), (10, "sic")):
                    yaz(ws, r, c, round(m["T"][k], 1), fmt="#,##0.0")
                yaz(ws, r, 9, m["T"]["eds"] / IL[dev]["Tk"], fmt="0.0%")
            yaz(ws, r, 11, ac)
            yaz(ws, r, 12, S["aciklik"][ac], fmt="0.0")
            for c, k in ((13, "buz"), (14, "sic"), (15, "min"), (16, "eds")):
                yaz(ws, r, c, round(sh[k], 3), fmt="0.00")
            r += 1
r += 1
ws.cell(row=r, column=1, value="İLETKEN VERİLERİ").font = Font(name="Arial", bold=True, size=11)
r += 1
bas = ["Devre", "İletken", "S (mm²)", "d (mm)", "G (kg/m)", "Kopma (kg)", "E (kg/mm²)", "α (1/°C)", "c",
       "Buz Pb=0,3√d (kg/m)", "G+buz (kg/m)", "Rüzgâr c·p·d (kg/m)", "Rüzgârlı bileşke (kg/m)", "Tmax (kg)",
       "σmax (kg/mm²)", "Tmax/kopma"]
for c, h in enumerate(bas, 1):
    cell = ws.cell(row=r, column=c, value=h)
    cell.font, cell.fill, cell.alignment, cell.border = BAS_FONT, BAS_FILL, ORTA, CER
ws.row_dimensions[r].height = 32
r += 1
for dev in ("OG", "AG"):
    il = IL[dev]
    vals = [dev, il["ad"], il["S"], il["d"], il["w0"], il["Tk"], il["E"], il["alfa"], il["c"]]
    for c, v in enumerate(vals, 1):
        yaz(ws, r, c, v, fmt="0.0000000" if c == 8 else None, al=SOL if c == 2 else None)
    yaz(ws, r, 10, f"=0.3*SQRT(D{r})", fmt="0.0000")
    yaz(ws, r, 11, f"=E{r}+J{r}", fmt="0.0000")
    yaz(ws, r, 12, f"=I{r}*44*D{r}/1000", fmt="0.0000")
    yaz(ws, r, 13, f"=SQRT(E{r}^2+L{r}^2)", fmt="0.0000")
    yaz(ws, r, 14, il["Tmax"], fmt="0.0")
    yaz(ws, r, 15, f"=N{r}/C{r}", fmt="0.00")
    yaz(ws, r, 16, f"=N{r}/F{r}", fmt="0.0%")
    r += 1
r += 1
for t_ in [("Profil kontrolleri (en büyük salgıda): OG yerden en az " +
           f"{min(v['OG_min_yukseklik'] for v in S['profil'].values()):.2f} m (≥ 7,0 m), AG en az "
           f"{min(v['AG_min_yukseklik'] for v in S['profil'].values() if 'AG_min_yukseklik' in v):.2f} m (≥ 5,0 m); "
           f"OG-AG en küçük düşey ayrım {min(v['OG_AG_min_ayrim'] for v in S['profil'].values() if 'OG_AG_min_ayrim' in v):.2f} m (≥ 1,50 m).").replace(".", ",").rstrip(",") + ".",
           f"Faz aralığı: D ≥ 0,5·√(fmax) + U/150 = {S['faz_araligi_min']:.3f} m ; potansta seçilen en küçük faz aralığı 0,90 m.".replace("0.7", "0,7")]:
    ws.cell(row=r, column=1, value=t_).font = F
    r += 1
sayfa_ayari(ws)

# ============================================================ TEMEL
ws = wb.create_sheet("TEMEL-BETON")
ust_bilgi(ws, "TEMEL HESABI (SULZBERGER) VE BETON / KAZI METRAJI",
          "Ms = b·t³·Ct·tanα/36 ; Mb = G·a·(0,5-0,47·√(G/(b·a²·Ct·tanα))) ; Ct = C2·t/2 ; Md = F·(H+2t/3) ; "
          "şart: Ms+Mb ≥ k·Md. Sarı hücreler değiştirilebilir (zemin katsayısı, başlık yüksekliği).", 20)
ws["A3"] = "C(2 m) zemin katsayısı (kg/cm³):"
ws["E3"] = S["sabitler"]["C2"]
ws["E3"].fill = GIR_FILL
ws["F3"] = "tanα:"
ws["G3"] = 0.01
ws["G3"].fill = GIR_FILL
ws["H3"] = "Beton yoğ. (kg/m³):"
ws["J3"] = 2200
ws["J3"].fill = GIR_FILL
ws["K3"] = "Yağmurluk başlığı (m):"
ws["M3"] = 0.20
ws["M3"].fill = GIR_FILL
ws["N3"] = "H (m):"
ws["O3"] = 9.55
for c in ("A3", "F3", "H3", "K3", "N3"):
    ws[c].font = FB
bas = ["Direk", "Tip", "b (m) kuvvete dik", "a (m) kuvvet yönü", "t (m)", "Kazı (m³)", "Beton (m³)", "Kalıp (m²)",
       "Direk ağırlığı (kg)", "G toplam (kg)", "Ct (kg/m³)", "Ms (kg·m)", "Mb (kg·m)", "Ms/Mb", "k",
       "F tepe (kg)", "Md (kg·m)", "(Ms+Mb)/(k·Md)", "Sonuç", "Katalog tepe kuvvetiyle GK"]
baslik_satiri(ws, 5, bas, [7, 8, 10, 10, 7, 10, 10, 10, 11, 11, 13, 12, 12, 8, 8, 10, 12, 13, 10, 13])
r = 6
for d in DIR:
    t = T[d]
    yaz(ws, r, 1, d, font=FB)
    yaz(ws, r, 2, t["tip"])
    yaz(ws, r, 3, t["b"], fmt="0.00")
    yaz(ws, r, 4, t["a"], fmt="0.00")
    yaz(ws, r, 5, t["t"], fmt="0.00")
    yaz(ws, r, 6, f"=C{r}*D{r}*E{r}", fmt="0.000")
    yaz(ws, r, 7, f"=C{r}*D{r}*(E{r}+$M$3)", fmt="0.000", font=FB)
    yaz(ws, r, 8, f"=2*(C{r}+D{r})*0.30", fmt="0.00")
    yaz(ws, r, 9, D[d]["agirlik"], fmt="#,##0.00")
    yaz(ws, r, 10, f"=G{r}*$J$3+I{r}+150", fmt="#,##0")
    yaz(ws, r, 11, f"=$E$3*E{r}/2*1000000", fmt="#,##0")
    yaz(ws, r, 12, f"=C{r}*E{r}^3*K{r}*$G$3/36", fmt="#,##0")
    yaz(ws, r, 13, f"=J{r}*D{r}*(0.5-0.47*SQRT(J{r}/(C{r}*D{r}^2*K{r}*$G$3)))", fmt="#,##0")
    yaz(ws, r, 14, f"=L{r}/M{r}", fmt="0.00")
    yaz(ws, r, 15, f"=IF(N{r}>=1,1,VLOOKUP(N{r},$W$6:$X$16,2,TRUE))", fmt="0.000")
    yaz(ws, r, 16, round(D[d]["F_tepe"], 1), fmt="#,##0.0")
    yaz(ws, r, 17, f"=P{r}*($O$3+2*E{r}/3)", fmt="#,##0")
    yaz(ws, r, 18, f"=(L{r}+M{r})/(O{r}*Q{r})", fmt="0.00", font=FB)
    yaz(ws, r, 19, f'=IF(R{r}>=1,"UYGUN","BÜYÜT")', font=FB)
    yaz(ws, r, 20, f"=(L{r}+M{r})/(O{r}*{D[d]['kapasite']}*($O$3+2*E{r}/3))", fmt="0.00")
    r += 1
# Sulzberger k tablosu
yaz(ws, 5, 23, "Ms/Mb", font=BAS_FONT, fill=BAS_FILL)
yaz(ws, 5, 24, "k", font=BAS_FONT, fill=BAS_FILL)
for i, (x_, k_) in enumerate([(0.0, 1.5), (0.1, 1.383), (0.2, 1.317), (0.3, 1.260), (0.4, 1.208), (0.5, 1.150),
                              (0.6, 1.115), (0.7, 1.075), (0.8, 1.040), (0.9, 1.017), (1.0, 1.0)]):
    yaz(ws, 6 + i, 23, x_, fmt="0.0")
    yaz(ws, 6 + i, 24, k_, fmt="0.000")
yaz(ws, r, 5, "TOPLAM", font=FB, fill=TOP_FILL)
for c in (6, 7, 8, 9):
    L = get_column_letter(c)
    yaz(ws, r, c, f"=SUM({L}6:{L}{r - 1})", fmt="#,##0.000" if c < 9 else "#,##0.00", font=FB, fill=TOP_FILL)
r += 2
for t_ in [
    "k katsayısı: W:X sütunlarındaki Sulzberger tablosundan (alt basamak, emniyetli); tüm direklerde Ms/Mb > 1 olduğundan k = 1,0.",
    "Temel boyutları TEDAŞ müşterek demir direk tip temelleridir (çift üs, III. bölge). C(2m)=8 kg/cm³ orta sert zemin kabulüdür; "
    "zemin etüdüne göre E3 hücresini güncelleyiniz.",
    "Beton C20/25 (TS EN 206), monoblok, demirsiz; direk ayakları temel tabanından 10 cm yukarıda biter. Üst yüz zeminden 20 cm "
    "yüksek ve eğimli yapılır.",
    "Katalog tepe kuvvetiyle GK < 1 çıkan K5'' temelleri, direk tam kapasitesinde zemin kabulüne duyarlıdır; gerçek yükte (GK sütunu R) "
    "yeterlidir. Zayıf zeminde temel boyutu büyütülmelidir.",
]:
    ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=20)
    ws.cell(row=r, column=1, value=t_).font = F
    r += 1
sayfa_ayari(ws)

# ======================================================= PROJE ESASLARI
ws = wb.create_sheet("PROJE ESASLARI")
ust_bilgi(ws, "PROJE ESASLARI VE KABULLER", "Hesaplarda kullanılan tüm giriş değerleri ve dayanakları.", 3)
baslik_satiri(ws, 4, ["Konu", "Değer / kabul", "Dayanak"], [28, 110, 34])
pe = [
    ("Şebeke", "34,5 kV OG (Um 36 kV), 3 faz, 50 Hz + 0,4 kV AG; müşterek (OG+AG aynı direk)", "Proje kabulü"),
    ("Buz yükü bölgesi", "III. bölge: Pb = 0,3·√d (kg/m), buz -5 °C, en düşük -25 °C, en yüksek +40 °C; buz yoğunluğu 0,6 kg/dm³",
     "EKATY Md.45, Çizelge-9"),
    ("Rüzgâr basıncı", "İletkenler 44 kg/m² (0-15 m, küçük aralıklı); direk/travers/izolatör 55 kg/m²",
     "EKATY Md.48, Çizelge-11"),
    ("Rüzgâr katsayısı c", "Pigeon (12,5<d≤15,8 mm) 1,1 ; AER (d>15,8) 1,0 ; kare kafes direk 2,8 ; tek yüzlü kafes/travers 1,6",
     "EKATY Çizelge-10"),
    ("OG iletken", "3/0 AWG Pigeon ACSR 6/1: S 99,30 mm², d 12,75 mm, 0,3439 kg/m, kopma 2.995 kg, E 8.000 kg/mm², α 19,1e-6",
     "EMO teknik bilgiler / TS"),
    ("OG çekme", "σmax = 5,0 kg/mm² → Tmax = 496,5 kg (%16,6 kopma ≤ %45); EDS ≈ %4,3 (≤ %15)", "EKATY Md.46"),
    ("AG kablo", "AER (Alpek) 3x95+70 mm²: demet d≈44 mm, 1,30 kg/m (emniyetli kabul), taşıyıcı nötr 70 mm² kopma 2.100 kg",
     "TS 11654 / TEDAŞ-MYD/2005-051 (teyit)"),
    ("AG çekme", "Tmax = 650 kg (-5 °C + buz), %31 kopma; EDS ≈ %13,5", "EKATY Md.46"),
    ("Direkler", "TEDAŞ müşterek galvanizli kafes demir direkler, a=50 m (çift üs ''), III. bölge; toprak üstü 9,73 m",
     "TEDAŞ tip direk tabloları"),
    ("Tepe kuvveti", "Tüm kuvvetler OG bağlantı seviyesine (9,55 m) indirgenmiş; AG 8,05 m (D1'de 7,05 m)",
     "EKATY Md.3 tanım 13"),
    ("Yüklenme varsayımları", "Taşıyıcı V1-V3, köşe durdurucu V1, V2, V4, V5; nihayet V1, V2, V4; Çizelge-12 zayıflama (4 iletken → %60)",
     "EKATY Md.49"),
    ("İzolatörler", "36 kV silikon kompozit: gergi 70 kN (durdurucu/köşe/nihayet), mesnet line-post (taşıyıcı ve atlama desteği)",
     "TEDAŞ şartnamesi"),
    ("Ayırıcı", "D1: 36 kV 630 A, harici tip, 3 kutuplu, silikon izolatörlü, kilitli kumanda kolu, potansiyel düzenleme",
     "TEDAŞ şartnamesi"),
    ("Mesafeler", "OG-AG direkte ≥ 1,50 m; yerden OG ≥ 7,0 m, AG yalıtılmış ≥ 5,0 m; faz aralığı ≥ 0,746 m",
     "EKATY Md.44, Çizelge-8"),
    ("Temel", "Monoblok C20/25, Sulzberger; C(2 m) = 8 kg/cm³; beton 2.200 kg/m³", "EKATY Md.56, Çizelge-19"),
    ("Topraklama", "Her direkte koruma topraklaması; D1 halka + 4 çubuk + kumanda yeri ızgarası; D2 ve D6'da AG nötr",
     "Elektrik Tesislerinde Topraklamalar Yönetmeliği"),
    ("Güzergâh", "M-D1 35 m, D1-D2 45 m, D2-D3 50 m, D3-D4 50 m, D4-D5 50 m, D5-D6 45 m; iç açılar D1 70°, D2 80°, D5 40°",
     "Proje kabulü (aplikasyonla teyit)"),
]
r = 5
for a, b, c in pe:
    yaz(ws, r, 1, a, font=FB, al=SOL)
    yaz(ws, r, 2, b, al=SOL)
    yaz(ws, r, 3, c, al=SOL)
    ws.row_dimensions[r].height = 28
    r += 1
sayfa_ayari(ws)

wb.save(OUT)
print("XLSX:", OUT)
